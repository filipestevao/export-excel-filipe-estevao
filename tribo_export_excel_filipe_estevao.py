# ==============================================================================
# Script Name: Export Excel Tribo by Filipe Estevao
# Description: Python scripts for use with Anton Paar's software for tribometer
# data export to Excel
# 
# Copyright (c) 2026 Filipe Estevão
# 
# Permission is hereby granted, free of charge, to any person obtaining a copy
# of this software and associated documentation files (the "Software"), to deal
# in the Software without restriction, including without limitation the rights
# to use, copy, modify, merge, publish, distribute, sublicense, and/or sell
# copies of the Software, and to permit persons to whom the Software is
# furnished to do so, subject to the following conditions:
# 
# The above copyright notice and this permission notice shall be included in all
# copies or substantial portions of the Software.
# 
# THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND, EXPRESS OR
# IMPLIED, INCLUDING BUT NOT LIMITED TO THE WARRANTIES OF MERCHANTABILITY,
# FITNESS FOR A PARTICULAR PURPOSE AND NONINFRINGEMENT. IN NO EVENT SHALL THE
# AUTHORS OR COPYRIGHT HOLDERS BE LIABLE FOR ANY CLAIM, DAMAGES OR OTHER
# LIABILITY, WHETHER IN AN ACTION OF CONTRACT, TORT OR OTHERWISE, ARISING FROM,
# OUT OF OR IN CONNECTION WITH THE SOFTWARE OR THE USE OR OTHER DEALINGS IN THE
# SOFTWARE.
# ==============================================================================

__title__ = "Export Excel Tribo by Filipe Estevao"
__version__ = "1.1.2"
__author__ = "Filipe Estevao"
__status__ = "Production"
__url__ = "https://github.com/filipestevao/export-excel-filipe-estevao"

import json
import math
import os
import subprocess
import sys
import urllib.request

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import BarChart, ScatterChart, Reference, Series
from openpyxl.chart.legend import Legend, LegendEntry
from openpyxl.chart.data_source import (
    AxDataSource,
    NumData,
    NumDataSource,
    NumRef,
    NumVal,
    StrData,
    StrRef,
    StrVal,
)
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Font, PatternFill, numbers
from openpyxl.utils.cell import get_column_letter, quote_sheetname

from antonpaar import connect_to_tribo
from antonpaar.script_tools import info


CURVE_TYPES = ['ctCombinated', 'ctStatic', 'ctDataCurves', 'ctFreeCurves']
STAT_ROWS = ['Mean', 'Std dev']
X_NAMES = ['Time']
Y_NAMES = ['µ', 'μ', 'Mu', 'Friction coefficient', 'Ft']
LAPS_NAMES = ['Laps', 'Lap']
DIST_NAMES = ['Distance', 'Sliding distance']
LAPS_INDEX = 3
DIST_INDEX = 4
EXTRA_RESULT_COLUMNS = (
    (
        'Lab temperature',
        '°C',
        'sequence.lab_environment.lab_temperature',
    ),
    (
        'Humidity',
        '%',
        'sequence.lab_environment.humidity',
    ),
)
SUMMARY_CHART_NAMES = ('Mu', 'Sample wear rate')
GROUP_COLORS = (
    '2B78C4', 'B0689C', '53A6A6', 'FF9DA7', 'FF9B00',
    '71C840', 'E15759', '9C755F', '8C8C8C', '28292E')
CURVE_LINE_WIDTH = 19050  # 1.5 pt in EMU (1 pt = 12700 EMU)
MAX_CURVE_POINTS = 5000


def is_number(value):
    return isinstance(value, (int, float)) and not math.isnan(value)


def stats(values):
    values = [v for v in values if is_number(v)]
    if not values:
        return {'Mean': None, 'Std dev': None, 'Min': None, 'Max': None, 'N': 0}
    mean = sum(values) / len(values)
    if len(values) > 1:
        variance = sum((v - mean) ** 2 for v in values) / (len(values) - 1)
        std_dev = math.sqrt(variance)
    else:
        std_dev = 0
    return {
        'Mean': mean,
        'Std dev': std_dev,
        'Min': min(values),
        'Max': max(values),
        'N': len(values),
    }


def curve_dimension(curves, curve_type, names):
    if curve_type not in curves:
        return None
    names = [n.lower() for n in names]
    curve_items = sorted(
        curves[curve_type].items(),
        key=lambda item: int(item[0]),
    )
    # Earlier names take precedence over later ones, regardless of
    # dimension index (e.g. prefer Mu over Ft even if Ft comes first).
    for name in names:
        for _, meta in curve_items:
            display = meta.get('DisplayName', '').lower()
            short = meta.get('ShortName', '').lower()
            if display == name or short == name:
                return int(meta.get('DimIndex'))
    for name in names:
        for _, meta in curve_items:
            display = meta.get('DisplayName', '').lower()
            short = meta.get('ShortName', '').lower()
            if name in display or name in short:
                return int(meta.get('DimIndex'))
    return None


def curve_unit(curves, curve_type, dim_index):
    meta = curves[curve_type].get(str(dim_index), {})
    unit = meta.get('PhysicalUnit', {})
    return unit.get('Symbol', ''), unit.get('SICoef', 1.0) or 1.0


def curve_name(curves, curve_type, dim_index):
    meta = curves[curve_type].get(str(dim_index), {})
    return (
        meta.get('ShortName')
        or meta.get('DisplayName')
        or 'Dim %d' % dim_index
    )


def curve_dim_count(curves, curve_type):
    try:
        return max(int(k) for k in curves[curve_type]) + 1
    except (ValueError, TypeError):
        return 0


def curve_dim_or_index(curves, curve_type, names, index):
    dim = curve_dimension(curves, curve_type, names)
    if dim is None:
        count = curve_dim_count(curves, curve_type)
        if 0 <= index < count:
            return index
    return dim


def curve_display_name(curves, curve_type, dim_index, fallback):
    meta = curves[curve_type].get(str(dim_index), {})
    return (
        meta.get('ShortName')
        or meta.get('DisplayName')
        or fallback
    )


def selected_acquisitions(groups):
    selected = []
    for group_id in groups['indexes']:
        group = groups['groups'][group_id]
        acquisitions = []
        for acquisition_index, data_id in enumerate(group['indexes'], 1):
            acquisition = group['data'][data_id]
            if acquisition.get('relevant'):
                acquisitions.append((data_id, acquisition, acquisition_index))
        if acquisitions:
            selected.append((group_id, group, acquisitions))
    return selected


def get_curve_data(server, doc_id, data_id, curve_type):
    first = server.curves.getdata(
        doc_id=doc_id,
        data_id=data_id,
        page_index=0,
        page_size=1,
        curve_type=curve_type)
    count = first.get('count', 0)
    if count == 0:
        return []
    data = server.curves.getdata(
        doc_id=doc_id,
        data_id=data_id,
        page_index=0,
        page_size=count,
        curve_type=curve_type)
    return data.get('data', [])


def get_breakpoints(server, doc_id, data_id, count):
    try:
        breakpoints = server.curves.breakpoints(
            doc_id=doc_id,
            data_id=data_id,
            curviline=True).get('breakpoints', [])
    except Exception:
        breakpoints = []
    breakpoints = [int(i) for i in breakpoints if 0 <= int(i) < count]
    if 0 not in breakpoints:
        breakpoints.insert(0, 0)
    if count and count - 1 not in breakpoints:
        breakpoints.append(count - 1)
    return sorted(set(breakpoints))


def resample_curve_indices(count, breakpoints, max_points):
    indices = list(range(count))
    if count <= max_points or len(breakpoints) < 2:
        return indices, breakpoints
    stride = math.ceil(count / max_points)
    if stride <= 1:
        return indices, breakpoints
    new_indices = []
    new_breakpoints = []
    for seg in range(len(breakpoints) - 1):
        start = breakpoints[seg]
        stop = breakpoints[seg + 1]
        seg_indices = list(range(start, stop + 1, stride))
        if seg_indices[-1] != stop:
            seg_indices.append(stop)
        new_breakpoints.append(len(new_indices))
        new_indices.extend(seg_indices)
    new_breakpoints.append(len(new_indices) - 1)
    return new_indices, new_breakpoints


def measurement_name(group, acquisition, acquisition_index):
    name = acquisition.get('name') or 'Measurement %d' % acquisition_index
    return '%s - %s' % (group['name'], name)


def curve_measurement_name(group, acquisition, acquisition_index):
    name = acquisition.get('name') or str(acquisition_index)
    return '%s - %s' % (group['name'], name)


def summary_chart_name(parameter):
    compact = parameter.upper().replace(' ', '')
    if (
        'MU' in compact
        or 'FRICTION' in compact
        or 'µ' in parameter
        or 'μ' in parameter
    ):
        return 'Mu'
    if 'WEAR' in compact:
        if 'PARTNER' in compact:
            return None
        return 'Sample wear rate'
    return None


def tribo_y_label(y_name):
    lower = y_name.lower()
    if (
        'mu' in lower
        or 'µ' in lower
        or 'μ' in lower
        or 'coefficient' in lower
    ):
        return 'μ (Coefficient of Friction)'
    return y_name


def fill_mu_from_curves(summary_chart_data, exported, y_name, y_unit):
    lower = y_name.lower()
    if (
        'mu' not in lower
        and 'µ' not in lower
        and 'μ' not in lower
        and 'coefficient' not in lower
    ):
        return
    values_by_group = {}
    for item in exported:
        points = [
            v for v in item['y'].tolist() if is_number(v)
        ]
        if not points:
            continue
        mean = sum(points) / len(points)
        values_by_group.setdefault(item['group'], []).append(mean)
    curve_groups = []
    for group_name, values in values_by_group.items():
        result = stats(values)
        if result['Mean'] is None:
            continue
        curve_groups.append({
            'group': group_name,
            'mean': result['Mean'],
            'std_dev': result['Std dev'],
        })
    if curve_groups:
        summary_chart_data['Mu'] = {
            'unit': y_unit,
            'groups': curve_groups,
        }


def color_for_group(group_name, group_colors):
    if group_name not in group_colors:
        color_index = len(group_colors)
        group_colors[group_name] = varied_color(
            GROUP_COLORS[color_index % len(GROUP_COLORS)],
            color_index // len(GROUP_COLORS))
    return group_colors[group_name]


def varied_color(hex_color, variation):
    if variation == 0:
        return hex_color

    values = [int(hex_color[i:i + 2], 16) for i in range(0, 6, 2)]
    if variation % 2:
        factor = max(0.35, 1.0 - 0.15 * ((variation + 1) // 2))
        values = [int(value * factor) for value in values]
    else:
        factor = min(0.55, 0.15 * (variation // 2))
        values = [int(value + (255 - value) * factor) for value in values]
    return ''.join('%02X' % value for value in values)


def parameter_value(
    server,
    doc_id,
    data_id,
    param_id,
    unit_factor,
    cycle_index=None,
    cycle_key='cycle_index',
):
    kwargs = {
        'doc_id': doc_id,
        'data_id': data_id,
        'param_id': param_id,
    }
    if cycle_index is not None:
        kwargs[cycle_key] = cycle_index
    value = server.parameters.getvalue(**kwargs)
    if isinstance(value, dict):
        if value.get('defined', True) is False:
            return None
        value = value.get('value')
    if is_number(value):
        return value / unit_factor
    return None


def safe_acquisition_analyses(server, doc_id, data_id):
    try:
        return server.acquisitions.analyses(
            doc_id=doc_id,
            acquisition_id=data_id).get('result', [])
    except Exception as error:
        info(' - skipping analyses for acquisition %s: %s' % (data_id, error))
        return []


def safe_doc_parameters(server, doc_id):
    try:
        params = server.parameters(doc_id=doc_id)
    except Exception as error:
        info(' - skipping document parameters: %s' % error)
        return {}
    if isinstance(params, dict):
        return params
    return {}


def acquisition_result_value(
    server,
    doc_id,
    acquisition_id,
    analysis_id,
    param_id,
    unit_factor,
):
    for data_id in (analysis_id, acquisition_id):
        try:
            value = parameter_value(
                server,
                doc_id,
                data_id,
                param_id,
                unit_factor,
            )
        except Exception:
            continue
        if value is not None:
            return value
    return None




def progress_average_segment(segment_x, segment_y, target_count):
    progress = np.linspace(0.0, 1.0, target_count)
    x_grid = []
    y_grid = []
    for x, y in zip(segment_x, segment_y):
        source_progress = np.linspace(0.0, 1.0, len(x))
        x_grid.append(np.interp(progress, source_progress, x))
        y_grid.append(np.interp(progress, source_progress, y))
    y_data = np.array(y_grid)
    if len(y_grid) > 1:
        y_std = np.std(y_data, axis=0, ddof=1)
    else:
        y_std = np.zeros(target_count)
    return (
        np.mean(np.array(x_grid), axis=0),
        np.mean(y_data, axis=0),
        y_std,
    )


def average_segment(segment_x, segment_y, target_count):
    return progress_average_segment(segment_x, segment_y, target_count)


def condition_path(conditions, path):
    value = conditions
    for part in path.split('.'):
        if not isinstance(value, dict) or part not in value:
            return None
        value = value[part]
    return value if is_number(value) else None


def write_results_sheet(
    wb, server, doc_id, selected, analyses_classes,
    server_version=None,
):
    ws = wb.create_sheet('Results')
    ws.freeze_panes = 'B1'
    header_fill = PatternFill('solid', fgColor='FCE4D6')
    section_fill = PatternFill('solid', fgColor='D9EAF7')
    summary_chart_data = {
        name: {'unit': '', 'groups': []}
        for name in SUMMARY_CHART_NAMES
    }
    row = 1

    ws.cell(row, 1, 'Tribometer results')
    ws.cell(row, 1).font = Font(bold=True)
    row += 2

    conditions_supported = (
        server_version is None
        or not _version_greater('11.0.0', server_version)
    )

    for _, group, acquisitions in selected:
        ws.cell(row, 1, group['name'])
        ws.cell(row, 1).font = Font(bold=True)
        ws.cell(row, 1).fill = section_fill
        row += 1

        columns = []
        columns_by_key = {}
        acquisition_names = []
        for data_id, acquisition, acquisition_index in acquisitions:
            acquisition_name = measurement_name(
                group,
                acquisition,
                acquisition_index,
            )
            acquisition_names.append(acquisition_name)
            analyses = safe_acquisition_analyses(server, doc_id, data_id)
            for analysis in analyses:
                analysis_class = analyses_classes.get(analysis['class_id'])
                if not analysis_class:
                    continue
                for parameter in analysis_class.get('parameters', []):
                    unit_factor = parameter.get('unit_factor', 1.0) or 1.0
                    value = acquisition_result_value(
                        server,
                        doc_id,
                        data_id,
                        analysis['id'],
                        parameter['id'],
                        unit_factor)
                    if value is not None:
                        key = (
                            analysis_class.get('class_name', ''),
                            parameter.get('name', ''),
                            parameter.get('unit', ''),
                        )
                        if key not in columns_by_key:
                            columns_by_key[key] = {
                                'analysis': analysis_class.get(
                                    'class_name',
                                    '',
                                ),
                                'parameter': parameter.get('name', ''),
                                'unit': parameter.get('unit', ''),
                                'values': [],
                                'acquisitions': {},
                            }
                            columns.append(columns_by_key[key])
                        columns_by_key[key]['values'].append(value)
                        columns_by_key[key]['acquisitions'][
                            acquisition_name
                        ] = value

        if not columns:
            doc_params = safe_doc_parameters(server, doc_id)
            for data_id, acquisition, acquisition_index in acquisitions:
                acquisition_name = measurement_name(
                    group,
                    acquisition,
                    acquisition_index,
                )
                for param_id, meta in doc_params.items():
                    if not isinstance(meta, dict):
                        continue
                    name = meta.get('name', '')
                    if name.startswith('General data'):
                        name = name[len('General data'):].lstrip(
                            ' -–—:')
                        if not name:
                            continue
                    unit_factor = meta.get('unit_factor', 1.0) or 1.0
                    try:
                        value = parameter_value(
                            server,
                            doc_id,
                            data_id,
                            param_id,
                            unit_factor,
                        )
                    except Exception:
                        continue
                    if value is not None:
                        key = (
                            'General data',
                            name,
                            meta.get('unit', ''),
                        )
                        if key not in columns_by_key:
                            columns_by_key[key] = {
                                'analysis': 'General data',
                                'parameter': name,
                                'unit': meta.get('unit', ''),
                                'values': [],
                                'acquisitions': {},
                            }
                            columns.append(columns_by_key[key])
                        columns_by_key[key]['values'].append(value)
                        columns_by_key[key]['acquisitions'][
                            acquisition_name
                        ] = value

        if conditions_supported:
            for data_id, acquisition, acquisition_index in acquisitions:
                acquisition_name = measurement_name(
                    group,
                    acquisition,
                    acquisition_index,
                )
                try:
                    cond = server.acquisitions.conditions(
                        doc_id=doc_id, acquisition_id=data_id,
                    )
                except Exception as exc:
                    if _is_method_unavailable(exc):
                        conditions_supported = False
                        break
                    continue
                if not isinstance(cond, dict):
                    continue
                for parameter, unit, path in EXTRA_RESULT_COLUMNS:
                    value = condition_path(cond, path)
                    if value is None:
                        continue
                    key = ('General data', parameter, unit)
                    if key not in columns_by_key:
                        columns_by_key[key] = {
                            'analysis': 'General data',
                            'parameter': parameter,
                            'unit': unit,
                            'values': [],
                            'acquisitions': {},
                        }
                        columns.append(columns_by_key[key])
                    columns_by_key[key]['values'].append(value)
                    columns_by_key[key]['acquisitions'][
                        acquisition_name
                    ] = value

        for column in columns:
            column['stats'] = stats(column['values'])

        group_chart_names = set()
        matched = {}
        for column in columns:
            chart_name = summary_chart_name(column['parameter'])
            if not chart_name or chart_name in group_chart_names:
                continue
            matched.setdefault(chart_name, []).append(column)
        for chart_name, chart_columns in matched.items():
            group_chart_names.add(chart_name)
            column = next(
                (
                    item for item in chart_columns
                    if 'MEAN' in item['parameter'].upper()
                ),
                chart_columns[0],
            )
            if not summary_chart_data[chart_name]['unit']:
                summary_chart_data[chart_name]['unit'] = column['unit']
            summary_chart_data[chart_name]['groups'].append({
                'group': group['name'],
                'mean': column['stats']['Mean'],
                'std_dev': column['stats']['Std dev'],
            })

        if not columns:
            ws.cell(
                row,
                1,
                'No analysis results found for selected measurements',
            )
            row += 2
            continue

        ws.cell(row, 1, '')
        for col, column in enumerate(columns, 2):
            title = column['parameter']
            if column['unit']:
                title += ' [%s]' % column['unit']
            ws.cell(row, col, title)
            ws.cell(row, col).fill = header_fill
            ws.cell(row, col).font = Font(bold=True)

        for offset, stat_name in enumerate(STAT_ROWS, 1):
            ws.cell(row + offset, 1, stat_name)
            ws.cell(row + offset, 1).font = Font(bold=True)
            for col, column in enumerate(columns, 2):
                ws.cell(row + offset, col, column['stats'][stat_name])

        row += len(STAT_ROWS) + 1
        for acquisition_name in acquisition_names:
            ws.cell(row, 1, acquisition_name)
            for col, column in enumerate(columns, 2):
                ws.cell(row, col, column['acquisitions'].get(acquisition_name))
            row += 1

        row += 2

    # Auto-size first column A to its content. openpyxl can't perfectly
    # match Excel AutoFit, so compute measured width and apply padding.
    # Build list of longest line lengths for each non-empty cell.
    lengths = []
    for cell in ws['A']:
        if cell.value is None:
            continue
        text = str(cell.value)
        # splitlines() returns empty list for empty string; fall back to
        # original text to ensure there's at least one item.
        lines = text.splitlines() or [text]
        lengths.append(max(len(line) for line in lines))
    max_len = max(lengths) if lengths else 16
    width = math.ceil(max_len * 1.05) + 2
    width = min(max(width, 8), 255)
    col_dim = ws.column_dimensions[get_column_letter(1)]
    col_dim.width = width

    for col in range(2, min(ws.max_column, 24) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16

    return summary_chart_data


def choose_curve(curves):
    for curve_type in CURVE_TYPES:
        x_dim = curve_dimension(curves, curve_type, X_NAMES)
        y_dim = curve_dimension(curves, curve_type, Y_NAMES)
        if x_dim is not None and y_dim is not None:
            laps_dim = curve_dim_or_index(
                curves, curve_type, LAPS_NAMES, LAPS_INDEX)
            dist_dim = curve_dim_or_index(
                curves, curve_type, DIST_NAMES, DIST_INDEX)
            return curve_type, x_dim, y_dim, laps_dim, dist_dim
    raise RuntimeError(
        'Could not find Time and Mu/Ft curve dimensions in the document'
    )


def write_curves_sheet(
    wb,
    server,
    doc_id,
    selected,
    curves,
    curve_type,
    x_dim,
    y_dim,
    laps_dim,
    dist_dim,
    group_colors,
):
    ws = wb.create_sheet('Curves')
    ws.freeze_panes = 'A3'
    x_unit, x_factor = curve_unit(curves, curve_type, x_dim)
    y_unit, y_factor = curve_unit(curves, curve_type, y_dim)
    x_name = curve_name(curves, curve_type, x_dim)
    y_name = curve_name(curves, curve_type, y_dim)
    y_label = tribo_y_label(y_name)
    specs = [(x_dim, x_factor)]
    headers = [chart_title(x_name, x_unit)]
    if laps_dim is not None:
        laps_unit, laps_factor = curve_unit(
            curves, curve_type, laps_dim)
        laps_name = curve_display_name(
            curves, curve_type, laps_dim, 'Laps')
        specs.append((laps_dim, laps_factor))
        headers.append(chart_title(laps_name, laps_unit))
    if dist_dim is not None:
        dist_unit, dist_factor = curve_unit(
            curves, curve_type, dist_dim)
        dist_name = curve_display_name(
            curves, curve_type, dist_dim, 'Distance')
        specs.append((dist_dim, dist_factor))
        # Metadata leaves Distance unitless; values are SI metres.
        headers.append(chart_title(dist_name, dist_unit or 'm'))
    specs.append((y_dim, y_factor))
    headers.append(chart_title(y_label, y_unit))
    exported = []
    chart = None
    col = 1

    for _, group, acquisitions in selected:
        for data_id, acquisition, acquisition_index in acquisitions:
            rows = get_curve_data(server, doc_id, data_id, curve_type)
            if not rows:
                continue
            breakpoints = get_breakpoints(server, doc_id, data_id, len(rows))
            indices, breakpoints = resample_curve_indices(
                len(rows), breakpoints, MAX_CURVE_POINTS)
            rows = [rows[i] for i in indices]
            display_name = curve_measurement_name(
                group,
                acquisition,
                acquisition_index,
            )
            ws.cell(1, col, display_name)
            for offset, header in enumerate(headers):
                ws.cell(2, col + offset, header)
            x_values = []
            y_values = []
            for row_index, point in enumerate(rows, 3):
                values = [
                    point[dim] / factor for dim, factor in specs
                ]
                for offset, value in enumerate(values):
                    ws.cell(row_index, col + offset, value)
                if is_number(values[0]) and is_number(values[-1]):
                    x_values.append(values[0])
                    y_values.append(values[-1])
            exported.append({
                'name': display_name,
                'group': group['name'],
                'col': col,
                'y_offset': len(specs) - 1,
                'count': len(rows),
                'breakpoints': breakpoints,
                'x': np.array(x_values, dtype=float),
                'y': np.array(y_values, dtype=float),
            })
            col += len(specs) + 1

    if exported:
        chart = ScatterChart()
        chart.scatterStyle = 'lineMarker'
        chart.title = 'Tribometer curves'
        chart.x_axis.title = '%s [%s]' % (x_name, x_unit) if x_unit else x_name
        chart.y_axis.title = chart_title(y_label, y_unit)
        chart.x_axis.scaling.min = 0
        chart.y_axis.scaling.min = 0
        chart.width = 24
        chart.height = 14
        style_curve_chart(chart)
        for item in exported:
            x_values = Reference(
                ws,
                min_col=item['col'],
                min_row=3,
                max_row=item['count'] + 2,
            )
            y_values = Reference(
                ws,
                min_col=item['col'] + item['y_offset'],
                min_row=3,
                max_row=item['count'] + 2,
            )
            series = Series(y_values, x_values, title=item['name'])
            color = color_for_group(item['group'], group_colors)
            series.graphicalProperties.line.solidFill = color
            series.graphicalProperties.line.width = CURVE_LINE_WIDTH
            series.marker.graphicalProperties.solidFill = color
            series.marker.graphicalProperties.line.solidFill = color
            cache_num_ref(series.xVal, item['x'].tolist())
            cache_num_ref(series.yVal, item['y'].tolist())
            chart.series.append(series)

    return exported, x_name, x_unit, y_name, y_unit, chart


def write_average_curves_sheet(
    wb,
    exported,
    x_name,
    x_unit,
    y_name,
    y_unit,
    group_colors,
):
    ws = wb.create_sheet('Average curves')
    ws.freeze_panes = 'A3'
    groups = []
    for item in exported:
        if len(item['x']) < 2 or len(item['y']) < 2:
            continue
        if item['group'] not in groups:
            groups.append(item['group'])

    if not groups:
        ws.cell(2, 1, 'No curve data available for averaging')
        return None

    chart = ScatterChart()
    chart.scatterStyle = 'lineMarker'
    chart.title = 'Average tribometer curves'
    y_label = tribo_y_label(y_name)
    chart.x_axis.title = '%s [%s]' % (x_name, x_unit) if x_unit else x_name
    chart.y_axis.title = chart_title(y_label, y_unit)
    chart.x_axis.scaling.min = 0
    chart.y_axis.scaling.min = 0
    chart.width = 24
    chart.height = 14
    style_curve_chart(chart)
    col = 1

    for group_name in groups:
        usable_items = [
            item
            for item in exported
            if (
                item['group'] == group_name
                and len(item['x']) >= 2
                and len(item['y']) >= 2
            )
        ]
        if not usable_items:
            continue

        ws.cell(1, col, group_name)
        x_header = (
            'Mean %s [%s]' % (x_name, x_unit)
            if x_unit
            else 'Mean %s' % x_name
        )
        y_header = 'Mean ' + chart_title(y_label, y_unit)
        std_header = chart_title('μ StdDev', y_unit)
        for offset, header in enumerate(
            (x_header, y_header, x_header, y_header, std_header),
        ):
            ws.cell(2, col + offset, header)

        averaged_x = []
        averaged_y = []
        averaged_std = []
        segment_count = max(
            len(item['breakpoints']) - 1
            for item in usable_items
        )
        for segment_index in range(segment_count):
            segment_x = []
            segment_y = []
            target_count = 0
            for item in usable_items:
                if segment_index + 1 >= len(item['breakpoints']):
                    continue
                start = item['breakpoints'][segment_index]
                stop = item['breakpoints'][segment_index + 1]
                x = item['x'][start:stop + 1]
                y = item['y'][start:stop + 1]
                if segment_index == 0:
                    positive = np.where((x > 0) | (y > 0))[0]
                    if len(positive):
                        x = x[positive[0]:]
                        y = y[positive[0]:]
                    x = np.insert(x, 0, 0.0)
                    y = np.insert(y, 0, 0.0)
                if len(x) < 2 or len(y) < 2:
                    continue
                segment_x.append(x)
                segment_y.append(y)
                target_count = max(target_count, len(x))

            if target_count < 2 or not segment_x:
                continue

            x_mean, y_mean, y_std = average_segment(
                segment_x, segment_y, target_count)
            if averaged_x:
                x_mean = x_mean[1:]
                y_mean = y_mean[1:]
                y_std = y_std[1:]
            averaged_x.extend(x_mean.tolist())
            averaged_y.extend(y_mean.tolist())
            averaged_std.extend(y_std.tolist())

        if not averaged_x:
            continue

        for index, x_value in enumerate(averaged_x, 3):
            ws.cell(index, col, float(x_value))
            ws.cell(index, col + 1, float(averaged_y[index - 3]))

        last_row = len(averaged_x) + 2
        x_values = Reference(ws, min_col=col, min_row=3, max_row=last_row)
        y_values = Reference(ws, min_col=col + 1, min_row=3, max_row=last_row)
        series = Series(y_values, x_values, title=group_name)
        color = color_for_group(group_name, group_colors)
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = CURVE_LINE_WIDTH
        series.marker.graphicalProperties.solidFill = color
        series.marker.graphicalProperties.line.solidFill = color
        cache_num_ref(series.xVal, averaged_x)
        cache_num_ref(series.yVal, averaged_y)
        chart.series.append(series)
        std_count = min(50, len(averaged_x))
        std_idx = np.linspace(
            0, len(averaged_x) - 1, std_count).astype(int)
        for pos, flat in enumerate(std_idx, 3):
            ws.cell(pos, col + 2, float(averaged_x[flat]))
            ws.cell(pos, col + 3, float(averaged_y[flat]))
            ws.cell(pos, col + 4, float(averaged_std[flat]))
        std_last = 2 + std_count
        std_x = Reference(
            ws, min_col=col + 2, min_row=3, max_row=std_last)
        std_y = Reference(
            ws, min_col=col + 3, min_row=3, max_row=std_last)
        std_ref = range_ref(ws, col + 4, 3, std_last)
        std_vals = [averaged_std[i] for i in std_idx]
        std_series = Series(
            std_y, std_x, title='%s StdDev' % group_name)
        std_series.graphicalProperties.line.noFill = True
        std_series.marker.symbol = 'circle'
        std_series.marker.size = 5
        std_series.marker.graphicalProperties.noFill = True
        std_series.marker.graphicalProperties.line.noFill = True
        std_series.errBars = ErrorBars(
            errDir='y',
            errBarType='both',
            errValType='cust',
            plus=NumDataSource(
                numRef=NumRef(
                    f=std_ref, numCache=number_cache(std_vals))
            ),
            minus=NumDataSource(
                numRef=NumRef(
                    f=std_ref, numCache=number_cache(std_vals))
            ),
            spPr=GraphicalProperties(),
        )
        std_series.errBars.spPr.ln = LineProperties()
        std_series.errBars.spPr.ln.solidFill = color
        cache_num_ref(std_series.xVal, [
            averaged_x[i] for i in std_idx])
        cache_num_ref(std_series.yVal, [
            averaged_y[i] for i in std_idx])
        chart.series.append(std_series)
        if chart.legend is None:
            chart.legend = Legend()
        chart.legend.legendEntry.append(
            LegendEntry(idx=len(chart.series) - 1, delete=True))
        col += 6

    if chart.series:
        return chart

    return None


def chart_title(name, unit):
    return '%s [%s]' % (name, unit) if unit else name


def range_ref(ws, column, first_row, last_row):
    return '%s!$%s$%d:$%s$%d' % (
        quote_sheetname(ws.title),
        get_column_letter(column),
        first_row,
        get_column_letter(column),
        last_row)


def number_cache(values):
    points = [
        NumVal(idx=index, v=value)
        for index, value in enumerate(values)
        if is_number(value)
    ]
    return NumData(ptCount=len(values), pt=points)


def string_cache(values):
    points = [
        StrVal(idx=index, v='' if value is None else str(value))
        for index, value in enumerate(values)
    ]
    return StrData(ptCount=len(values), pt=points)


def cache_num_ref(data_source, values):
    if data_source and data_source.numRef:
        data_source.numRef.numCache = number_cache(values)


def set_chart_axis_ids(chart, base_id):
    x_axis_id = base_id + 1
    y_axis_id = base_id + 2
    chart.x_axis.axId = x_axis_id
    chart.y_axis.axId = y_axis_id
    chart.x_axis.crossAx = y_axis_id
    chart.y_axis.crossAx = x_axis_id
    chart.x_axis.axPos = 'b'
    chart.y_axis.axPos = 'l'
    for axis in (chart.x_axis, chart.y_axis):
        axis.delete = False
        axis.tickLblPos = 'nextTo'
        axis.crosses = 'autoZero'
        axis.numFmt = 'General'
    if isinstance(chart, ScatterChart):
        chart.x_axis.crossBetween = 'midCat'
        chart.y_axis.crossBetween = 'midCat'
    else:
        chart.y_axis.crossBetween = 'between'


def style_summary_bar_chart(chart):
    chart.y_axis.scaling.min = 0

    # Set overall chart border
    if chart.plot_area.spPr is None:
        chart.plot_area.spPr = GraphicalProperties()
    if chart.plot_area.spPr.ln is None:
        chart.plot_area.spPr.ln = LineProperties()
    chart.plot_area.spPr.ln.solidFill = '000000'

    chart.y_axis.majorGridlines.spPr = GraphicalProperties()
    chart.y_axis.majorGridlines.spPr.line.solidFill = 'C5C5C5'


def style_curve_chart(chart):
    # Set overall chart border
    if chart.plot_area.spPr is None:
        chart.plot_area.spPr = GraphicalProperties()
    if chart.plot_area.spPr.ln is None:
        chart.plot_area.spPr.ln = LineProperties()
    chart.plot_area.spPr.ln.solidFill = '000000'

    # Set internal gridlines (horizontal and vertical) to gray
    for axis in (chart.x_axis, chart.y_axis):
        axis.majorGridlines.spPr = GraphicalProperties()
        axis.majorGridlines.spPr.line.solidFill = 'C5C5C5'


def add_summary_bar_chart(
    ws,
    chart_name,
    chart_info,
    anchor,
    data_col,
    axis_base_id,
):
    groups = chart_info['groups']
    unit = chart_info['unit']
    if chart_name == 'Mu':
        title_name = 'μ (Coefficient of Friction)'
    else:
        title_name = chart_name

    ws.cell(2, data_col, title_name)
    ws.cell(3, data_col, 'Group')
    ws.cell(3, data_col + 1, 'Mean')
    ws.cell(3, data_col + 2, 'Std dev')

    for row_offset, item in enumerate(groups, 4):
        ws.cell(row_offset, data_col, item['group'])
        ws.cell(row_offset, data_col + 1, item['mean'])
        ws.cell(row_offset, data_col + 2, item['std_dev'])

    if not groups:
        ws.cell(4, data_col, 'No %s data found' % chart_name)
        return

    first_row = 4
    last_row = first_row + len(groups) - 1
    chart = BarChart()
    chart.title = chart_title(title_name, unit)
    chart.y_axis.title = chart_title(title_name, unit)
    chart.legend = None
    chart.width = 12
    chart.height = 8
    set_chart_axis_ids(chart, axis_base_id)
    style_summary_bar_chart(chart)
    if chart_name == 'Sample wear rate':
        chart.y_axis.numFmt = '0.00E+00'
    chart.type = 'col'
    chart.grouping = 'clustered'
    chart.add_data(
        Reference(ws, min_col=data_col + 1, min_row=3, max_row=last_row),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(
            ws,
            min_col=data_col,
            min_row=first_row,
            max_row=last_row,
        )
    )
    chart.series[0].cat = AxDataSource(strRef=StrRef(
        f=range_ref(ws, data_col, first_row, last_row),
        strCache=string_cache([item['group'] for item in groups])))
    cache_num_ref(chart.series[0].val, [item['mean'] for item in groups])

    std_values = [item['std_dev'] for item in groups]
    std_ref = range_ref(ws, data_col + 2, first_row, last_row)
    chart.series[0].errBars = ErrorBars(
        errDir='y',
        errBarType='both',
        errValType='cust',
        plus=NumDataSource(
            numRef=NumRef(f=std_ref, numCache=number_cache(std_values))
        ),
        minus=NumDataSource(
            numRef=NumRef(f=std_ref, numCache=number_cache(std_values))
        ),
    )
    ws.add_chart(chart, anchor)


def write_charts_sheet(
    ws,
    doc_name,
    summary_chart_data,
    curve_chart,
    average_curve_chart,
):
    ws.cell(1, 1, doc_name)
    ws.cell(1, 1).font = Font(bold=True)

    anchors = {'Mu': 'A3', 'Sample wear rate': 'I3'}
    data_col = 30
    axis_base_id = 1000
    for chart_name in SUMMARY_CHART_NAMES:
        add_summary_bar_chart(
            ws,
            chart_name,
            summary_chart_data[chart_name],
            anchors[chart_name],
            data_col,
            axis_base_id)
        data_col += 4
        axis_base_id += 100

    if curve_chart is not None:
        set_chart_axis_ids(curve_chart, axis_base_id)
        ws.add_chart(curve_chart, 'A20')
        axis_base_id += 100
    if average_curve_chart is not None:
        set_chart_axis_ids(average_curve_chart, axis_base_id)
        ws.add_chart(average_curve_chart, 'A48')


def _is_method_unavailable(exc):
    # JSON-RPC -32601 = method absent from server API, not a retriable error.
    if exc.args and exc.args[0] == -32601:
        return True
    return 'Method not found' in str(exc)


def format_condition_value(value):
    if isinstance(value, float):
        if value.is_integer():
            return '%d' % value
        return '%.6g' % value
    return str(value)


def flatten_conditions(value, prefix=''):
    lines = []
    if isinstance(value, dict):
        for key, item in value.items():
            path = '%s.%s' % (prefix, key) if prefix else str(key)
            lines.extend(flatten_conditions(item, path))
    elif isinstance(value, (list, tuple)):
        if all(
            not isinstance(item, (dict, list, tuple))
            for item in value
        ):
            texts = [format_condition_value(v) for v in value]
            if prefix:
                if texts:
                    lines.append('%s: %s' % (prefix, ', '.join(texts)))
            else:
                lines.extend(texts)
        else:
            for index, item in enumerate(value):
                lines.extend(flatten_conditions(
                    item, '%s[%d]' % (prefix, index)))
    elif value is None or value == '':
        pass
    else:
        text = format_condition_value(value)
        if prefix:
            lines.append('%s: %s' % (prefix, text))
        else:
            lines.append(text)
    return lines


def write_measurement_parameters_sheet(
    wb, server, doc_id, selected, server_version=None,
):
    ws = wb.create_sheet('Parameters')
    header_fill = PatternFill('solid', fgColor='D9EAF7')
    section_fill = PatternFill('solid', fgColor='FCE4D6')
    ws.cell(1, 1, 'Values as returned by the instrument (SI units)')
    ws.cell(1, 1).font = Font(italic=True)

    # Method absent before V11; calling it corrupts the TCP session.
    # Check version up front instead of retrying per measurement.
    conditions_supported = (
        server_version is None
        or not _version_greater('11.0.0', server_version)
    )

    measurements = []
    for _, group, acquisitions in selected:
        if not conditions_supported:
            break
        for data_id, acquisition, acquisition_index in acquisitions:
            name = measurement_name(
                group, acquisition, acquisition_index,
            )
            try:
                cond = server.acquisitions.conditions(
                    doc_id=doc_id, acquisition_id=data_id,
                )
            except Exception as exc:
                if _is_method_unavailable(exc):
                    # Stop retrying if method unexpectedly unavailable.
                    conditions_supported = False
                    break
                continue
            measurements.append({'name': name, 'conditions': cond})

    if not measurements:
        wb.remove(ws)
        return

    for m in measurements:
        cond = m['conditions']
        if isinstance(cond, dict):
            summary = cond.get('summary', [])
            indenter = cond.get('indenter', {})
            rest = {
                key: value for key, value in cond.items()
                if key not in ('summary', 'indenter')
            }
        else:
            summary = cond if isinstance(cond, list) else []
            indenter = {}
            rest = {}
        if not isinstance(indenter, dict):
            indenter = {}
        param_lines = list(summary) + flatten_conditions(rest)
        m['param_lines'] = param_lines
        m['indenter'] = indenter

        key_items = list(param_lines)
        for field in ('serial_number', 'material', 'geometry'):
            if field in indenter:
                key_items.append(str(indenter[field]))
        m['key'] = tuple(key_items)

    groups = {}
    for m in measurements:
        key = m['key']
        if key not in groups:
            groups[key] = {
                'measurements': [],
                'param_lines': m['param_lines'],
                'indenter': m['indenter'],
            }
        groups[key]['measurements'].append(m)

    def write_group(ws, group_data, col):
        row = 2
        ws.cell(row, col, 'Measurement Name').font = Font(bold=True)
        ws.cell(row, col).fill = header_fill
        row += 1
        for m in group_data['measurements']:
            ws.cell(row, col, m['name'])
            row += 1
        row += 1
        ws.cell(row, col, '# Tribometer Parameters').font = (
            Font(bold=True)
        )
        ws.cell(row, col).fill = section_fill
        row += 1
        for line in group_data['param_lines']:
            if line:
                ws.cell(row, col, line)
                row += 1
        row += 1
        indenter = group_data['indenter']
        if indenter:
            ws.cell(row, col, '# Indenters').font = Font(bold=True)
            ws.cell(row, col).fill = section_fill
            row += 1
            type_val = indenter.get('geometry', '')
            if type_val:
                ws.cell(row, col, 'Type : %s' % type_val)
                row += 1
            serial = indenter.get('serial_number', '')
            if serial:
                ws.cell(row, col, 'Serial number : %s' % serial)
                row += 1
            material = indenter.get('material', '')
            if material:
                ws.cell(row, col, 'Material : %s' % material)
                row += 1
        row += 2

    if len(groups) == 1:
        group_data = next(iter(groups.values()))
        write_group(ws, group_data, 1)
        ws.column_dimensions['A'].width = 55
    else:
        col = 1
        for group_data in groups.values():
            if col > 1:
                ws.column_dimensions[
                    get_column_letter(col)
                ].width = 10
                col += 1
            write_group(ws, group_data, col)
            ws.column_dimensions[
                get_column_letter(col)
            ].width = 55
            col += 1


def format_numbers(ws):
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, (int, float)):
                if cell.value != 0 and abs(cell.value) < 0.01:
                    cell.number_format = '0.00E+00'
                else:
                    cell.number_format = '#,##0.00'


def export_selected_tribo_excel(server, doc_id, server_version=None):
    docs = server.docs()
    doc = docs['docs'][doc_id]
    doc_path = doc.get('path') or doc.get('name') or 'tribo_export'
    export_path = os.path.splitext(doc_path)[0] + '.xlsx'
    groups = server.groups(doc_id=doc_id)
    selected = selected_acquisitions(groups)
    if not selected:
        raise RuntimeError(
            'No selected/relevant tribometer measurements found'
        )

    curves = server.curves(doc_id=doc_id)
    analyses_classes = {
        item['class_id']: item
        for item in server.analyses.classes().get('result', [])
    }
    curve_type, x_dim, y_dim, laps_dim, dist_dim = choose_curve(curves)

    info(
        'Exporting selected tribometer measurements from %s'
        % doc.get('name', doc_id)
    )
    info(' - curve type: %s' % curve_type)

    wb = Workbook(write_only=False)
    charts_ws = wb.active
    charts_ws.title = 'Charts'
    group_colors = {}
    summary_chart_data = write_results_sheet(
        wb,
        server,
        doc_id,
        selected,
        analyses_classes,
        server_version,
    )

    write_measurement_parameters_sheet(
        wb, server, doc_id, selected, server_version)

    exported, x_name, x_unit, y_name, y_unit, curve_chart = write_curves_sheet(
        wb,
        server,
        doc_id,
        selected,
        curves,
        curve_type,
        x_dim,
        y_dim,
        laps_dim,
        dist_dim,
        group_colors,
    )
    average_curve_chart = write_average_curves_sheet(
        wb,
        exported,
        x_name,
        x_unit,
        y_name,
        y_unit,
        group_colors,
    )
    fill_mu_from_curves(summary_chart_data, exported, y_name, y_unit)
    write_charts_sheet(
        charts_ws,
        os.path.basename(doc_path),
        summary_chart_data,
        curve_chart,
        average_curve_chart)

    for ws in wb.worksheets:
        format_numbers(ws)

    wb.save(export_path)
    info('File saved: %s' % export_path)
    return export_path


def open_file_with_default_program(filename):
    if sys.platform.startswith('win'):
        os.startfile(filename)
    elif sys.platform == 'darwin':
        subprocess.Popen(['open', filename])
    else:
        subprocess.Popen(['xdg-open', filename])


def ask_to_open_file(filename):
    import tkinter as tk
    from tkinter import messagebox

    root = tk.Tk()
    root.withdraw()
    root.attributes('-topmost', True)
    should_open = messagebox.askyesno(
        'Open generated file?',
        'The Excel file was created successfully.\n\nOpen it now?',
        parent=root)
    root.destroy()

    if should_open:
        open_file_with_default_program(filename)


def check_for_updates():
    info('%s - Current version: %s' % (__title__, __version__))
    try:
        repo_path = __url__.rstrip('/').replace(
            'https://github.com/', '')
        api_url = (
            'https://api.github.com/repos/%s/tags' % repo_path)
        req = urllib.request.Request(
            api_url,
            headers={'User-Agent': 'opencode/1.0'},
            method='GET')
        with urllib.request.urlopen(req, timeout=10) as resp:
            tags = json.loads(resp.read().decode())
    except Exception:
        info(
            'Could not check for updates: no internet connection\n'
        )
        return

    latest = ''
    for tag in tags:
        ver = tag['name'].lstrip('v')
        if _version_greater(ver, latest):
            latest = ver

    if not latest:
        info('Could not determine latest version\n')
        return

    try:
        [int(x) for x in __version__.split('.')]
        current_valid = True
    except (ValueError, AttributeError):
        current_valid = False

    if not current_valid or _version_greater(__version__, latest):
        info('Could not determine latest version\n')
        return

    if _version_greater(latest, __version__):
        info(
            'New version %s available. Visit %s to download\n'
            % (latest, __url__)
        )
    else:
        info('Latest version (%s) is installed\n' % __version__)


def _version_greater(a, b):
    if not a:
        return False
    if not b:
        return True
    try:
        parts_a = [int(x) for x in a.split('.')]
        parts_b = [int(x) for x in b.split('.')]
    except (ValueError, AttributeError):
        return False
    max_len = max(len(parts_a), len(parts_b))
    parts_a += [0] * (max_len - len(parts_a))
    parts_b += [0] * (max_len - len(parts_b))
    for pa, pb in zip(parts_a, parts_b):
        if pa != pb:
            return pa > pb
    return False


if __name__ == '__main__':
    check_for_updates()
    tribo = connect_to_tribo()
    result = tribo.ls()
    info('Connected to %(server_name)s, V%(server_version)s' % result)

    docs = tribo.docs()
    doc_id = docs.get('current') or docs['indexes'][0]
    filename = export_selected_tribo_excel(
        tribo, doc_id, result.get('server_version'))
    ask_to_open_file(filename)
